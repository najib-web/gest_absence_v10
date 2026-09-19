#!/usr/bin/env python3
"""Analyse la structure des fichiers templates (xlsx + csv)."""
import openpyxl, csv, sys

for path in ["/home/z/my-project/public/templates/ListEleve_20260905.xlsx",
             "/home/z/my-project/public/templates/Liste enseignants.xlsx"]:
    print(f"\n=== {path.split('/')[-1]} ===")
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        print(f"  Sheet: '{ws.title}' dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
        # print first 12 rows
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(12, ws.max_row), values_only=True)):
            print(f"    R{i+1}: {row}")

print("\n=== tableaux de services.csv ===")
with open("/home/z/my-project/public/templates/tableaux de services.csv", encoding="utf-8", errors="replace") as f:
    for i, line in enumerate(f):
        if i > 15: break
        print(f"  L{i+1}: {line.rstrip()}")
